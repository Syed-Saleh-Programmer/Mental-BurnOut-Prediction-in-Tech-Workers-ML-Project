import io

import pandas as pd
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.main import app
from backend.hr import TemporaryStore

DATA = 'ID,Name,Department,Role,Stress,Burnout,Sleep,Weekly hours\na,Ana,Engineering,Developer,8,7,5,50\nb,Ben,Support,Analyst,4,3,7.5,40\nc,Cam,Engineering,Developer,,5.5,6.5,45\nd,Dee,Support,Analyst,oops,8,6,48\n'


@pytest.fixture(scope='module')
def client():
    with TestClient(app) as instance:
        yield instance


def upload(client, content=DATA, filename='team.csv'):
    response = client.post('/api/hr/imports', files={'file': (filename, content)})
    assert response.status_code == 200, response.text
    return response.json()


def commit(client, content=DATA):
    draft = upload(client, content)
    body = {'mapping': draft['mapping'], 'accept_valid_rows': True}
    response = client.post(f"/api/hr/imports/{draft['import_id']}/analyze", json=body)
    assert response.status_code == 200, response.text
    return response.json()


def test_import_preview_suggests_mapping_without_changing_source(client):
    draft = upload(client)
    assert draft['row_count'] == 4
    assert draft['mapping']['ID'] == 'employee_id'
    assert draft['mapping']['Weekly hours'] == 'work_hours_per_week'
    assert client.get('/api/health').json()['dataset_rows'] == 100000


def test_invalid_rows_require_explicit_acceptance_and_downloadable_issues(client):
    draft = upload(client)
    url = f"/api/hr/imports/{draft['import_id']}"
    body = {'mapping': draft['mapping']}
    check = client.post(url + '/validate', json=body).json()
    assert (check['valid_rows'], check['rejected_rows'], check['eligible_predictions']) == (3, 1, 0)
    assert client.post(url + '/analyze', json=body).status_code == 422
    report = client.get(url + '/issues.csv')
    assert 'Stress' in report.text and '5,Stress' in report.text
    assert client.post(url + '/analyze', json={**body, 'accept_valid_rows': True}).status_code == 200


def test_upload_metrics_and_filtered_counts_use_only_accepted_rows(client):
    analysis = commit(client)
    base = f"/api/hr/analyses/{analysis['analysis_id']}"
    overview = client.get(base + '/overview').json()
    metrics = {m['key']: m for m in overview['metrics']}
    assert overview['count'] == 3
    assert metrics['stress_score']['value'] == 6
    assert metrics['stress_score']['n'] == 2
    assert metrics['burnout_score']['value'] == 5.17
    filtered = client.get(base + '/overview?department=Engineering').json()
    assert filtered['count'] == 2
    assert next(m for m in filtered['metrics'] if m['key'] == 'stress_score')['value'] == 8
    exported = pd.read_csv(io.StringIO(client.get(base + '/exports/employees.csv?department=Engineering').text))
    assert len(exported) == 2
    assert set(exported.department) == {'Engineering'}
    assert 'Reported stress' in client.get(base + '/exports/summary.csv').text


def test_employee_search_sort_profile_and_scope(client):
    analysis = commit(client)
    base = f"/api/hr/analyses/{analysis['analysis_id']}"
    data = client.get(base + '/employees?search=Cam').json()
    assert data['count'] == 1
    row = data['records'][0]
    assert row['employee_name'] == 'Cam' and row['stress_score'] is None
    assert client.get(base + '/employees/' + row['row_id']).json()['employee_name'] == 'Cam'
    assert client.get(base + '/employees/row-999').status_code == 404
    sorted_rows = client.get(base + '/employees?sort_by=stress_score&direction=desc&page_size=2').json()['records']
    assert [r['employee_name'] for r in sorted_rows] == ['Ana', 'Ben']
    assert client.get(base + '/employees?sort_by=made_up').status_code == 422


def test_complete_profiles_keep_original_model_probabilities(client):
    source = app.state.service.df.head(2).drop(columns='job_change_intention')
    expected = app.state.service.predict(source.to_dict('records'))['predictions']
    analysis = commit(client, source.to_csv(index=False))
    assert analysis['eligible_predictions'] == 2
    rows = client.get(f"/api/hr/analyses/{analysis['analysis_id']}/employees").json()['records']
    assert [r['prediction']['probability'] for r in rows] == [p['probability'] for p in expected]


def test_cbi_score_is_not_converted_into_source_burnout(client):
    draft = upload(client, 'ID,Burnout\na,80\nb,40\n')
    spec = {'mapping': draft['mapping'], 'burnout_scale': 'cbi100'}
    response = client.post(f"/api/hr/imports/{draft['import_id']}/analyze", json=spec)
    assert response.status_code == 200
    overview = client.get(f"/api/hr/analyses/{response.json()['analysis_id']}/overview").json()
    assert overview['metrics'] == [{'key': 'cbi_score', 'label': 'CBI exhaustion', 'unit': '/ 100', 'value': 60, 'n': 2}]
    assert overview['distribution_kind'] == 'cbi_score'


def test_excel_selects_a_usable_sheet_and_allows_switching(client):
    book = Workbook(); book.active.title = 'Empty'
    first = book.create_sheet('First team'); first.append(['ID', 'Stress']); first.append(['a', 8])
    second = book.create_sheet('Second team'); second.append(['ID', 'Stress']); second.append(['b', 2])
    stream = io.BytesIO(); book.save(stream)
    draft = upload(client, stream.getvalue(), 'teams.xlsx')
    assert draft['sheet'] == 'First team'
    switched = client.get(f"/api/hr/imports/{draft['import_id']}/preview?sheet=Second%20team").json()
    assert switched['preview'][0]['Stress'] == 2
    result = client.post(f"/api/hr/imports/{draft['import_id']}/analyze", json={'mapping': switched['mapping'], 'sheet': 'Second team'})
    assert result.status_code == 200
    metric = client.get(f"/api/hr/analyses/{result.json()['analysis_id']}/overview").json()['metrics'][0]
    assert metric['value'] == 2


@pytest.mark.parametrize('name,content', [('bad.xls', 'x'), ('bad.xlsx', b'not a workbook'), ('empty.csv', ''), ('duplicate.csv', 'Stress,Stress\n3,5'), ('bad.csv', b'\xff\xfe'), ('ragged.csv', 'ID,Stress\na,4,extra')])
def test_unreadable_files_have_actionable_errors(client, name, content):
    response = client.post('/api/hr/imports', files={'file': (name, content)})
    assert response.status_code == 422
    assert 'detail' in response.json()


def test_upload_limits(client):
    assert client.post('/api/hr/imports', files={'file': ('big.csv', b'x' * (10*1024*1024+1))}).status_code == 413
    assert client.post('/api/hr/imports', files={'file': ('many.csv', 'Stress\n' + '4\n' * 5001)}).status_code == 422


def test_mapping_and_duplicate_identifiers_are_reviewed(client):
    draft = upload(client, 'ID,Stress,Sleep\na,4,6\na,5,7\n')
    base = f"/api/hr/imports/{draft['import_id']}"
    result = client.post(base + '/validate', json={'mapping': draft['mapping']}).json()
    assert result['rejected_rows'] == 2
    assert any('Duplicate' in issue['message'] for issue in result['issues'])
    assert client.post(base + '/validate', json={'mapping': {'ID': 'employee_id'}}).status_code == 422
    assert client.post(base + '/validate', json={'mapping': {'Stress': 'stress_score', 'Sleep': 'stress_score'}}).status_code == 422


def test_decimal_comma_is_not_silently_misread(client):
    draft = upload(client, 'ID,Sleep\na,"0,5"\n')
    result = client.post(f"/api/hr/imports/{draft['import_id']}/validate", json={'mapping': draft['mapping']}).json()
    assert result['rejected_rows'] == 1


def test_clear_and_missing_context_never_fall_back_to_research(client):
    analysis = commit(client)
    base = f"/api/hr/analyses/{analysis['analysis_id']}"
    assert client.delete(base).status_code == 200
    assert client.get(base + '/overview').status_code == 410
    assert client.get('/api/hr/analyses/unknown/overview').status_code == 410


def test_store_expiry_refresh_and_bounded_capacity():
    now = [0]
    store = TemporaryStore(ttl=5, capacity=2, clock=lambda: now[0])
    a = store.put('analysis', {'one': 1}); now[0] = 4
    store.get(a, 'analysis'); now[0] = 6
    assert store.get(a, 'analysis')['one'] == 1
    b = store.put('analysis', {}); store.put('analysis', {})
    with pytest.raises(HTTPException): store.get(a, 'analysis')
    now[0] = 20; store.cleanup(); assert not store.items


def test_export_escapes_formula_prefixes(client):
    analysis = commit(client, 'ID,Name,Stress\na,=1+1,4\n')
    text = client.get(f"/api/hr/analyses/{analysis['analysis_id']}/exports/employees.csv").text
    assert "'=1+1" in text


def test_frontend_deep_links_and_unknown_api_status(client):
    assert client.get('/hr/employees/row-2').status_code == 200
    assert client.get('/individual/explore').status_code == 200
    assert client.get('/api/does-not-exist').status_code == 404


def test_scenario_reports_targets_without_changing_or_predicting_burnout(client):
    result = client.post('/api/self-assessment/scenario', json={'current': {'sleep_hours': 5.5, 'work_hours': 52, 'stress': 0}, 'scenario': {'sleep_hours': 7, 'work_hours': 40, 'stress': 2}})
    assert result.status_code == 200
    data = result.json(); rows = {r['id']: r for r in data['comparisons']}
    assert rows['sleep_hours']['delta'] == 1.5
    assert rows['work_hours']['delta'] == -12
    assert rows['stress']['current'] == 0
    assert rows['manager_support']['delta'] is None
    assert 'burnout' not in data and 'probability' not in data
    assert client.post('/api/self-assessment/scenario', json={'current': {}, 'scenario': {'sleep_hours': 25}}).status_code == 422
