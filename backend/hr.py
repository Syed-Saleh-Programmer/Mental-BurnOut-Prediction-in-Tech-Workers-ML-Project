"""Temporary, upload-scoped people analytics. Never changes the research source."""
from __future__ import annotations

import csv
import io
import json
import math
import re
import secrets
import threading
import time
import zipfile
from collections import Counter, OrderedDict
from typing import Annotated, Any, Literal

import numpy as np
import pandas as pd
from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field

router = APIRouter(prefix='/api/hr', tags=['HR workspace'])
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
GROUPS = ['department', 'team', 'job_role', 'work_mode', 'country']
METRICS = {
    'burnout_score': ('Reported burnout', '/ 10'), 'cbi_score': ('CBI exhaustion', '/ 100'),
    'stress_score': ('Reported stress', '/ 10'), 'sleep_hours_per_night': ('Nightly sleep', 'hrs / night'),
    'work_hours_per_week': ('Working hours', 'hrs / week'), 'work_life_balance_score': ('Work-life balance', '/ 10'),
    'manager_support_score': ('Manager support', '/ 10'), 'job_satisfaction_score': ('Job satisfaction', '/ 10'),
}
DOMAIN = {
    'age': (0, 120), 'years_experience': (0, 100), 'years_at_company': (0, 100),
    'salary_usd': (0, 1e9), 'team_size': (0, 100000), 'work_hours_per_week': (0, 168),
    'sleep_hours_per_night': (0, 24), 'meetings_per_day': (0, 100), 'vacation_days_taken': (0, 366),
    'cbi_score': (0, 100), 'stress_score': (0, 10), 'burnout_score': (0, 10),
}
ALIASES = {
    'id': 'employee_id', 'employeenumber': 'employee_id', 'staffid': 'employee_id',
    'name': 'employee_name', 'fullname': 'employee_name', 'employeename': 'employee_name',
    'dept': 'department', 'division': 'department', 'jobtitle': 'job_role', 'role': 'job_role',
    'location': 'country', 'workingmode': 'work_mode', 'stress': 'stress_score',
    'burnout': 'burnout_score', 'cbiscore': 'cbi_score', 'cbiburnoutscore': 'cbi_score',
    'sleep': 'sleep_hours_per_night', 'sleephours': 'sleep_hours_per_night',
    'weeklyhours': 'work_hours_per_week', 'workhours': 'work_hours_per_week', 'hoursperweek': 'work_hours_per_week',
    'balance': 'work_life_balance_score', 'worklifebalance': 'work_life_balance_score',
    'support': 'manager_support_score', 'managersupport': 'manager_support_score',
    'satisfaction': 'job_satisfaction_score', 'jobsatisfaction': 'job_satisfaction_score',
    'burnoutcategory': 'burnout_level',
}


class TemporaryStore:
    """Bounded RAM cache, refreshed on use. References are not authentication."""
    def __init__(self, ttl=1800, capacity=12, clock=time.monotonic):
        self.ttl, self.capacity, self.clock = ttl, capacity, clock
        self.items: OrderedDict[str, dict] = OrderedDict()
        self.lock = threading.RLock()

    def cleanup(self):
        with self.lock:
            for key in list(self.items):
                if self.items[key]['expires'] <= self.clock():
                    del self.items[key]

    def put(self, kind: str, value: dict) -> str:
        with self.lock:
            self.cleanup()
            while len(self.items) >= self.capacity:
                self.items.popitem(last=False)
            key = f'{kind}_{secrets.token_urlsafe(20)}'
            self.items[key] = {'kind': kind, 'value': value, 'expires': self.clock() + self.ttl}
            return key

    def get(self, key: str, kind: str) -> dict:
        with self.lock:
            self.cleanup()
            entry = self.items.get(key)
            if not entry or entry['kind'] != kind:
                raise HTTPException(410, 'This temporary workspace has expired or was cleared. Upload your file again.')
            entry['expires'] = self.clock() + self.ttl
            self.items.move_to_end(key)
            return entry['value']

    def delete(self, key: str):
        with self.lock:
            self.items.pop(key, None)


def services(request):
    return request.app.state.service, request.app.state.hr_store


def destinations(project):
    fields = [{'key': 'employee_name', 'label': 'Employee name', 'kind': 'text'},
              {'key': 'department', 'label': 'Department', 'kind': 'text'},
              {'key': 'team', 'label': 'Team', 'kind': 'text'},
              {'key': 'cbi_score', 'label': 'CBI score (0–100)', 'kind': 'number'},
              {'key': 'job_change_intention', 'label': 'Reported intention to leave (0/1)', 'kind': 'integer'}]
    fields += [{'key': f['name'], 'label': f['label'], 'kind': f['kind']} for f in project.fields]
    return fields


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', str(value).lower())


def suggested_mapping(headers, project):
    aliases = {normalized(f['key']): f['key'] for f in destinations(project)} | ALIASES
    used, mapping = set(), {}
    for header in headers:
        key = aliases.get(normalized(header), '')
        mapping[header] = key if key not in used else ''
        if key:
            used.add(key)
    return mapping


def read_table(draft, sheet=None):
    if draft['kind'] == 'csv':
        try:
            reader = csv.reader(io.StringIO(draft['bytes'].decode('utf-8-sig')), strict=True)
            headers = next(reader, [])
            rows, positions = [], []
            for number, row in enumerate(reader, start=2):
                if number > 100000:
                    raise HTTPException(422, 'Too many worksheet rows. Remove empty rows and upload again.')
                if not any(str(v).strip() for v in row):
                    continue
                if len(row) != len(headers):
                    raise HTTPException(422, f'CSV row {number} has a different number of cells than its header.')
                rows.append(row); positions.append(number)
                if len(rows) > MAX_ROWS:
                    raise HTTPException(422, 'Upload at most 5,000 employee rows.')
        except (UnicodeError, csv.Error) as exc:
            raise HTTPException(422, 'Use a UTF-8 comma-separated CSV with one header row.') from exc
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(draft['bytes']), read_only=True, data_only=True)
        try:
            selected = sheet or draft['sheets'][0]
            if selected not in draft['sheets']:
                raise HTTPException(422, 'Choose one of the available worksheets.')
            worksheet = workbook[selected]
            if worksheet.max_column and worksheet.max_column > 200:
                raise HTTPException(422, 'Use a worksheet with at most 200 columns.')
            iterator = worksheet.iter_rows(values_only=True)
            headers = list(next(iterator, []))
            while headers and headers[-1] is None:
                headers.pop()
            rows, positions = [], []
            for number, row in enumerate(iterator, start=2):
                if number > 50002:
                    raise HTTPException(422, 'Remove excess empty/formatting rows and upload the worksheet again.')
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                rows.append(list(row[:len(headers)])); positions.append(number)
                if len(rows) > MAX_ROWS:
                    raise HTTPException(422, 'Upload at most 5,000 employee rows.')
        finally:
            workbook.close()
    headers = [str(h).strip() if h is not None else '' for h in headers]
    if not headers or not rows:
        raise HTTPException(422, 'The selected file or worksheet has no employee records.')
    if len(headers) > 200 or any(not h or len(h) > 128 for h in headers):
        raise HTTPException(422, 'Use 1–200 named columns, with headers no longer than 128 characters.')
    if len(set(h.casefold() for h in headers)) != len(headers):
        raise HTTPException(422, 'Column headers must be unique. Rename duplicate columns and upload again.')
    return pd.DataFrame(rows, columns=headers, index=positions)


def json_records(frame):
    # Pandas serializes missing numeric values as JSON null rather than NaN.
    return json.loads(frame.to_json(orient='records', date_format='iso', double_precision=10))


def preview(draft, project, import_id, sheet=None):
    table = read_table(draft, sheet)
    return {'import_id': import_id, 'filename': draft['filename'], 'sheets': draft['sheets'],
            'sheet': sheet or (draft['sheets'][0] if draft['sheets'] else None),
            'headers': table.columns.tolist(), 'row_count': len(table), 'preview': json_records(table.head(20)),
            'mapping': suggested_mapping(table.columns, project), 'fields': destinations(project)}


class MappingRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    mapping: dict[str, str] = Field(max_length=200)
    sheet: str | None = Field(default=None, max_length=128)
    burnout_scale: Literal['source10', 'cbi100'] = 'source10'
    accept_valid_rows: bool = False


def blank(value):
    return value is None or (isinstance(value, float) and math.isnan(value)) or str(value).strip() == ''


def validate_table(table, spec, project):
    allowed = {f['key'] for f in destinations(project)}
    mapping = {key: value for key, value in spec.mapping.items() if value}
    if any(key not in table.columns or value not in allowed for key, value in mapping.items()):
        raise HTTPException(422, 'The column mapping contains an unknown source or destination.')
    if len(set(mapping.values())) != len(mapping):
        raise HTTPException(422, 'Map each destination field only once.')
    if spec.burnout_scale == 'cbi100' and 'burnout_score' in mapping.values() and 'cbi_score' in mapping.values():
        raise HTTPException(422, 'Map a 0–100 score to CBI only once.')
    if not set(mapping.values()) & (set(METRICS) | {'burnout_level'}):
        raise HTTPException(422, 'Map at least one wellbeing indicator, such as stress, burnout, sleep, or working hours.')
    numeric = {f['name']: f for f in project.fields if f['kind'] != 'category' and f['name'] != 'employee_id'}
    numeric['job_change_intention'] = {'kind': 'integer', 'min': 0, 'max': 1}
    categories = {f['name']: {v.casefold(): v for v in f['options']} for f in project.fields if f['kind'] == 'category'}
    ids = Counter(str(row[source]).strip() for _, row in table.iterrows() for source, target in mapping.items() if target == 'employee_id' and not blank(row[source]))
    valid, issues = [], []
    for position, row in table.iterrows():
        record, errors = {'row_id': f'row-{position}', 'source_row': int(position)}, []
        for source, target in mapping.items():
            value = row[source]
            if spec.burnout_scale == 'cbi100' and target == 'burnout_score':
                target = 'cbi_score'
            if blank(value):
                record[target] = None
                if target == 'employee_id':
                    errors.append({'field': source, 'message': 'A mapped employee ID cannot be blank.', 'value': ''})
                continue
            text = str(value).strip()
            if target in numeric or target == 'cbi_score':
                try:
                    binary = target in {'therapy_access', 'uses_therapy', 'ai_tools_daily', 'seeks_mental_health_support', 'job_change_intention'}
                    if binary and text.casefold() in ['yes', 'true', 'no', 'false']:
                        number = float(text.casefold() in ['yes', 'true'])
                    else:
                        cleaned = text
                        if target == 'salary_usd':
                            cleaned = cleaned.removeprefix('$').strip()
                        if re.fullmatch(r'[+-]?\d{1,3}(,\d{3})+(\.\d+)?', cleaned):
                            cleaned = cleaned.replace(',', '')
                        number = float(cleaned)
                    lo, hi = DOMAIN.get(target, (numeric.get(target, {}).get('min', 0), numeric.get(target, {}).get('max', 100)))
                    if target.endswith('_score') and target not in ['phq9_score', 'gad7_score', 'cbi_score']:
                        lo, hi = 0, 10
                    if not math.isfinite(number) or number < lo or number > hi:
                        raise ValueError(f'Use a finite value from {lo:g} to {hi:g}.')
                    if numeric.get(target, {}).get('kind') == 'integer' and number != int(number):
                        raise ValueError('Use a whole number.')
                    record[target] = int(number) if number == int(number) else number
                except (ValueError, OverflowError) as exc:
                    errors.append({'field': source, 'message': str(exc) if str(exc).startswith(('Use ', 'A ')) else 'Use a valid number.', 'value': text[:100]})
            else:
                if len(text) > 200:
                    errors.append({'field': source, 'message': 'Text must be 200 characters or fewer.', 'value': text[:100]})
                else:
                    record[target] = categories.get(target, {}).get(text.casefold(), text)
                if target == 'employee_id' and ids[text] > 1:
                    errors.append({'field': source, 'message': 'Duplicate employee ID. Use one row per employee.', 'value': text})
        if not any(record.get(key) is not None for key in [*METRICS, 'burnout_level']):
            errors.append({'field': 'Wellbeing indicators', 'message': 'This row has no usable wellbeing indicator.', 'value': ''})
        if errors:
            issues.extend({'row': int(position), **error} for error in errors)
            continue
        missing = [key for key in project.raw_columns if record.get(key) is None]
        model_input = None
        reason = f'Missing {len(missing)} model inputs.' if missing else ''
        if not missing:
            try:
                model_input = project.validate([{key: record[key] for key in project.raw_columns}])[0]
            except HTTPException:
                reason = 'One or more fields fall outside the saved model’s categories or training ranges.'
        record['model_eligible'] = model_input is not None
        record['model_missing_fields'] = missing
        record['model_reason'] = reason
        record['_model_input'] = model_input
        record['prediction'] = None
        record['prediction_probability'] = None
        valid.append(record)
    return valid, issues


def validation_response(table, valid, issues):
    rejected = len({i['row'] for i in issues})
    return {'total_rows': len(table), 'valid_rows': len(valid), 'rejected_rows': rejected,
            'eligible_predictions': sum(row['model_eligible'] for row in valid),
            'issues': issues[:100], 'issue_count': len(issues),
            'preview': [{k: v for k, v in row.items() if not k.startswith('_')} for row in valid[:5]]}


def create_analysis(valid, issues, project, store, filename, source_kind='upload', source_rows=None):
    eligible = [row for row in valid if row['model_eligible']]
    if eligible:
        results = project.predict([row['_model_input'] for row in eligible])['predictions']
        for row, prediction in zip(eligible, results):
            row['prediction'] = prediction
            row['prediction_probability'] = prediction['probability']
    for row in valid:
        row.pop('_model_input', None)
    # Immutable records; only cache expiry changes during reads.
    data = {'records': valid, 'issues': issues, 'filename': filename, 'source_kind': source_kind,
            'source_rows': source_rows or len(valid) + len({i['row'] for i in issues}), 'threshold': project.threshold}
    analysis_id = store.put('analysis', data)
    return metadata(analysis_id, data)


def metadata(analysis_id, data):
    return {'analysis_id': analysis_id, 'filename': data['filename'], 'source_kind': data['source_kind'],
            'row_count': len(data['records']), 'source_rows': data['source_rows'],
            'rejected_rows': len({i['row'] for i in data['issues']}),
            'eligible_predictions': sum(r['model_eligible'] for r in data['records']),
            'threshold': data['threshold'], 'expires_after_minutes': 30,
            'available_metrics': [key for key in METRICS if any(r.get(key) is not None for r in data['records'])]}


def selected_records(data, department=None, team=None, job_role=None, country=None, work_mode=None, prediction=None):
    result = data['records']
    for key, value in [('department', department), ('team', team), ('job_role', job_role), ('country', country), ('work_mode', work_mode)]:
        if value:
            result = [r for r in result if r.get(key) == value]
    if prediction in ['leave', 'stay']:
        result = [r for r in result if r.get('prediction') and r['prediction']['prediction'] == (1 if prediction == 'leave' else 0)]
    return result


def analytics(records, data, group_by=None):
    def metric(key):
        values = [r[key] for r in records if r.get(key) is not None]
        return {'key': key, 'label': METRICS[key][0], 'unit': METRICS[key][1],
                'value': round(float(np.mean(values)), 2) if values else None, 'n': len(values)}
    metrics = [metric(key) for key in METRICS if any(r.get(key) is not None for r in data['records'])]
    predictions = [r['prediction'] for r in records if r.get('prediction')]
    if any(r.get('prediction') for r in data['records']):
        metrics.append({'key': 'predicted_intention', 'label': 'Predicted intention to leave', 'unit': '%',
                        'value': round(float(np.mean([r['prediction'] for r in predictions])) * 100, 2) if predictions else None, 'n': len(predictions)})
    observed = [r['job_change_intention'] for r in records if r.get('job_change_intention') is not None]
    if any(r.get('job_change_intention') is not None for r in data['records']):
        metrics.append({'key': 'observed_intention', 'label': 'Reported intention to leave', 'unit': '%',
                        'value': round(float(np.mean(observed)) * 100, 2) if observed else None, 'n': len(observed)})
    group_keys = [key for key in GROUPS if any(r.get(key) for r in data['records'])]
    grouping = group_by if group_by in group_keys else next(iter(group_keys), None)
    primary = next((key for key in ['cbi_score', 'burnout_score', 'stress_score', 'work_hours_per_week', 'sleep_hours_per_night', 'job_satisfaction_score'] if any(r.get(key) is not None for r in data['records'])), None)
    groups = []
    if grouping:
        for name in sorted({r[grouping] for r in records if r.get(grouping)}):
            rows = [r for r in records if r.get(grouping) == name]
            values = [r[primary] for r in rows if primary and r.get(primary) is not None]
            groups.append({'name': name, 'count': len(rows), 'value': round(float(np.mean(values)), 2) if values else (len(rows) if primary is None else None), 'n': len(values) if primary else len(rows)})
        groups.sort(key=lambda g: -g['count'])
    score = next((key for key in ['cbi_score', 'burnout_score'] if any(r.get(key) is not None for r in data['records'])), None)
    categories = Counter(r['burnout_level'] for r in records if r.get('burnout_level'))
    distribution = []
    if categories:
        order = {'Low': 0, 'Moderate': 1, 'High': 2, 'Severe': 3}
        distribution = [{'name': k, 'count': v} for k, v in sorted(categories.items(), key=lambda pair: (order.get(pair[0], 4), pair[0]))]
    elif score:
        upper = 100 if score == 'cbi_score' else 10
        for i in range(5):
            lo, hi = i * upper / 5, (i + 1) * upper / 5
            distribution.append({'name': f'{lo:g}–{hi:g}', 'count': sum(r.get(score) is not None and lo <= r[score] <= hi and (i == 4 or r[score] < hi) for r in records)})
    relationship = []
    if score:
        for i in range(5):
            lo, hi = i * 2, i * 2 + 2
            rows = [r for r in records if r.get('stress_score') is not None and r.get(score) is not None and lo <= r['stress_score'] <= hi and (i == 4 or r['stress_score'] < hi)]
            if rows:
                relationship.append({'name': f'{lo}–{hi}', 'count': len(rows), 'value': round(float(np.mean([r[score] for r in rows])), 2)})
    insights = []
    for key in ['work_hours_per_week', 'sleep_hours_per_night']:
        m = next((m for m in metrics if m['key'] == key and m['n']), None)
        if m:
            insights.append(f"Average {m['label'].lower()} is {m['value']:g} {m['unit']}, based on {m['n']:,} responses.")
    if predictions:
        insights.append(f'{len(predictions):,} of {len(records):,} profiles have complete, compatible inputs for the job-change intention model.')
    if not insights:
        insights.append(f'{len(records):,} profiles match the selected cohort. Metrics use only the responses provided for each indicator.')
    return {'count': len(records), 'total_count': len(data['records']), 'metrics': metrics,
            'groups': groups[:12], 'group_by': grouping, 'group_keys': group_keys,
            'group_metric': METRICS.get(primary) if primary else None,
            'distribution': distribution, 'distribution_kind': 'reported_categories' if categories else score,
            'relationship': relationship, 'score_metric': METRICS.get(score) if score else None,
            'insights': insights[:3], 'eligible_predictions': len(predictions),
            'filters': {key: sorted({r[key] for r in data['records'] if r.get(key)}) for key in group_keys}}


def csv_download(records, filename):
    frame = pd.DataFrame(records)
    for col in frame.select_dtypes('object'):
        frame[col] = frame[col].map(lambda v: "'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@', '\t', '\r')) else v)
    return Response(frame.to_csv(index=False), media_type='text/csv', headers={'Content-Disposition': f'attachment; filename="{filename}"'})


@router.get('/schema')
def import_schema(request: Request):
    project, _ = services(request)
    return {'fields': destinations(project), 'max_rows': MAX_ROWS, 'max_bytes': MAX_BYTES}


@router.get('/template.csv')
def template(request: Request):
    project, _ = services(request)
    columns = ['employee_id', 'department', 'job_role', 'work_mode', 'stress_score', 'burnout_score', 'sleep_hours_per_night', 'work_hours_per_week', 'manager_support_score', 'job_satisfaction_score']
    sample = project.df.head(3).copy()
    sample['department'] = 'Example team'
    return csv_download(sample[columns].to_dict('records'), 'workwell-example-template.csv')


@router.post('/imports')
async def upload(request: Request, file: Annotated[UploadFile, File()]):
    project, store = services(request)
    name = file.filename or ''
    if not name.lower().endswith(('.csv', '.xlsx')):
        await file.close()
        raise HTTPException(422, 'Choose a CSV or .xlsx file. Convert older Excel files to .xlsx or CSV first.')
    try:
        content = await file.read(MAX_BYTES + 1)
    finally:
        await file.close()
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Upload a file of 10 MB or less.')
    draft = {'filename': name[:200], 'kind': 'csv' if name.lower().endswith('.csv') else 'xlsx', 'bytes': content, 'sheets': []}
    if draft['kind'] == 'xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(413, 'This workbook expands beyond the supported size. Export one worksheet as CSV.')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            draft['sheets'] = book.sheetnames
            book.close()
            if not draft['sheets']:
                raise ValueError('No worksheets')
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'This Excel file could not be read. Use an unencrypted .xlsx workbook.') from exc
    import_id = store.put('import', draft)
    try:
        if draft['sheets']:
            failures = []
            for sheet_name in draft['sheets']:
                try:
                    return preview(draft, project, import_id, sheet_name)
                except HTTPException as exc:
                    failures.append(exc)
            raise failures[0]
        return preview(draft, project, import_id)
    except Exception:
        store.delete(import_id)
        raise


@router.get('/imports/{import_id}/preview')
def sheet_preview(request: Request, import_id: str, sheet: str | None = None):
    project, store = services(request)
    return preview(store.get(import_id, 'import'), project, import_id, sheet)


@router.post('/imports/{import_id}/validate')
def validate_import(request: Request, import_id: str, body: MappingRequest):
    project, store = services(request)
    draft = store.get(import_id, 'import')
    table = read_table(draft, body.sheet)
    valid, issues = validate_table(table, body, project)
    draft['issues'] = issues
    return validation_response(table, valid, issues)


@router.get('/imports/{import_id}/issues.csv')
def import_issues(request: Request, import_id: str):
    _, store = services(request)
    return csv_download(store.get(import_id, 'import').get('issues', []), 'workwell-import-issues.csv')


@router.post('/imports/{import_id}/analyze')
def analyze_import(request: Request, import_id: str, body: MappingRequest):
    project, store = services(request)
    draft = store.get(import_id, 'import')
    table = read_table(draft, body.sheet)
    valid, issues = validate_table(table, body, project)
    if not valid:
        raise HTTPException(422, 'No valid employee rows are available. Review the issues and correct your file or mapping.')
    if issues and not body.accept_valid_rows:
        raise HTTPException(422, 'Review the rejected rows and explicitly choose to analyze the valid rows.')
    result = create_analysis(valid, issues, project, store, draft['filename'], source_rows=len(table))
    store.delete(import_id)
    return result


@router.post('/demo')
def demo(request: Request):
    project, store = services(request)
    table = project.df.head(MAX_ROWS).copy()
    table.index = range(2, len(table) + 2)
    spec = MappingRequest(mapping={key: key for key in project.df.columns})
    # Include the optional observed target in a demo/research-format upload.
    valid, issues = validate_table(table, spec, project)
    return create_analysis(valid, issues, project, store, 'Research dataset sample (5,000 profiles)', 'demo', len(project.df))


@router.get('/analyses/{analysis_id}')
def get_metadata(request: Request, analysis_id: str):
    _, store = services(request)
    return metadata(analysis_id, store.get(analysis_id, 'analysis'))


@router.get('/analyses/{analysis_id}/overview')
def overview(request: Request, analysis_id: str, department: str | None = None, team: str | None = None, job_role: str | None = None, country: str | None = None, work_mode: str | None = None, prediction: str | None = None, group_by: str | None = None):
    _, store = services(request)
    data = store.get(analysis_id, 'analysis')
    return analytics(selected_records(data, department, team, job_role, country, work_mode, prediction), data, group_by)


@router.get('/analyses/{analysis_id}/employees')
def employees(request: Request, analysis_id: str, department: str | None = None, team: str | None = None, job_role: str | None = None, country: str | None = None, work_mode: str | None = None, prediction: str | None = None,
              search: Annotated[str, Query(max_length=100)] = '', page: Annotated[int, Query(ge=1)] = 1, page_size: Annotated[int, Query(ge=1, le=100)] = 10,
              sort_by: str = 'source_row', direction: Literal['asc', 'desc'] = 'asc'):
    _, store = services(request)
    data = store.get(analysis_id, 'analysis')
    rows = selected_records(data, department, team, job_role, country, work_mode, prediction)
    if search:
        rows = [r for r in rows if any(search.casefold() in str(r.get(key, '')).casefold() for key in ['employee_id', 'employee_name', 'department', 'job_role', 'team', 'country'])]
    allowed = set(GROUPS) | set(METRICS) | {'employee_id', 'employee_name', 'source_row', 'prediction_probability', 'burnout_level'}
    if sort_by not in allowed:
        raise HTTPException(422, 'Choose a valid sort field.')
    # Empty values stay last in both directions.
    filled = [r for r in rows if r.get(sort_by) is not None]
    empty = [r for r in rows if r.get(sort_by) is None]
    filled.sort(key=lambda r: (r[sort_by] if isinstance(r[sort_by], (int, float)) else str(r[sort_by]).casefold()), reverse=direction == 'desc')
    rows = filled + empty
    return {'count': len(rows), 'page': page, 'page_size': page_size, 'records': rows[(page-1)*page_size:page*page_size]}


@router.get('/analyses/{analysis_id}/employees/{row_id}')
def employee(request: Request, analysis_id: str, row_id: str):
    _, store = services(request)
    row = next((r for r in store.get(analysis_id, 'analysis')['records'] if r['row_id'] == row_id), None)
    if row is None:
        raise HTTPException(404, 'Employee profile not found in this analysis.')
    return row


@router.get('/analyses/{analysis_id}/exports/{kind}.csv')
def export(request: Request, analysis_id: str, kind: Literal['summary', 'employees', 'issues'], department: str | None = None, team: str | None = None, job_role: str | None = None, country: str | None = None, work_mode: str | None = None, prediction: str | None = None):
    _, store = services(request)
    data = store.get(analysis_id, 'analysis')
    rows = selected_records(data, department, team, job_role, country, work_mode, prediction)
    if kind == 'issues':
        records = data['issues']
    elif kind == 'summary':
        summary = analytics(rows, data)
        records = [{'metric': 'Profiles analyzed', 'value': len(rows), 'n': len(rows), 'unit': 'profiles'}] + [{k: m[k] for k in ['label', 'value', 'n', 'unit']} for m in summary['metrics']]
        records[0]['label'] = records[0].pop('metric')
    else:
        records = [{**{k: v for k, v in r.items() if k not in ['prediction', 'model_missing_fields']}, 'predicted_intention': r['prediction']['label'] if r.get('prediction') else None} for r in rows]
    return csv_download(records, f'workwell-{kind}.csv')


@router.delete('/analyses/{analysis_id}')
def clear(request: Request, analysis_id: str):
    _, store = services(request)
    store.delete(analysis_id)
    return {'cleared': True}
